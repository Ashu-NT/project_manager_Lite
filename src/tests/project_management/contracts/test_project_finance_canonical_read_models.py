from __future__ import annotations

from dataclasses import replace
from datetime import date
from decimal import Decimal
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from src.core.modules.project_management.domain.financials.forecast import (
    ForecastGenerationMode,
    ForecastLineSourceKind,
    ForecastLineSourceType,
)
from src.core.modules.project_management.domain.financials.rate_cards import RateType
from src.core.modules.project_management.infrastructure.persistence.orm.baseline import (
    BaselineTaskORM,
)
from src.core.modules.project_management.infrastructure.persistence.orm.cost_entry import (
    ProjectCostEntryORM,
)
from src.core.modules.project_management.infrastructure.persistence.orm.forecast import (
    ForecastLineORM,
)
from src.core.modules.project_management.infrastructure.reporting import (
    api as reporting_api,
)
from src.core.platform.common.exceptions import NotFoundError


def _become_independent_decider(services) -> None:
    user_session = services["user_session"]
    user_session.set_principal(
        replace(
            user_session.principal,
            user_id="independent-cost-decider",
            username="independent-cost-decider",
        )
    )


def _login(services, username: str, password: str) -> None:
    auth = services["auth_service"]
    user = auth.authenticate(username, password)
    services["user_session"].set_principal(auth.build_principal(user))


def _approved_controls(services):
    organization = services["tenant_context_service"].get_active_organization()
    project = services["project_service"].create_project(
        "Canonical finance read model",
        financial_currency_code=organization.base_currency,
    )
    code = services["financial_configuration_service"].create_cost_code(
        code="D4-CONTROL",
        name="D.4 control",
    )

    budgets = services["budget_service"]
    budget = budgets.create_budget(project.id, "Approved budget")
    budgets.add_line(
        budget.id,
        cost_code_id=code.id,
        description="Authorized scope",
        amount=Decimal(100),
        expected_budget_version=budget.row_version,
    )
    budget = budgets.get_budget(budget.id)
    budget = budgets.submit_budget(
        budget.id, "admin", expected_version=budget.row_version
    )
    result = budgets.approve_budget(
        budget.id, approved_by="admin", expected_version=budget.row_version
    )
    budget = budgets.get_budget(result.budget_id)

    forecasts = services["forecast_version_service"]
    forecast = forecasts.create_forecast(
        project.id,
        name="Approved ETC",
        as_of_date=date(2026, 8, 1),
        generation_mode=ForecastGenerationMode.MANUAL,
        created_by="admin",
    )
    forecasts.add_line(
        forecast.id,
        cost_code_id=code.id,
        description="Remaining delivery",
        amount=Decimal(80),
        source_kind=ForecastLineSourceKind.MANUAL,
        source_type=ForecastLineSourceType.MANUAL_ESTIMATE,
        created_by="admin",
        expected_forecast_version=forecast.row_version,
        period_start=date(2026, 8, 1),
        period_end=date(2026, 8, 31),
    )
    forecast = forecasts.get_forecast(forecast.id)
    forecast = forecasts.submit_forecast(
        forecast.id,
        submitted_by="admin",
        expected_version=forecast.row_version,
    )
    forecast = forecasts.approve_forecast(
        forecast.id,
        approved_by="admin",
        expected_version=forecast.row_version,
    )

    services["financial_period_service"].create_period(
        code="D4-2026-08",
        name="August 2026",
        fiscal_year=2026,
        period_number=8,
        start_date=date(2026, 8, 1),
        end_date=date(2026, 8, 31),
    )
    entries = services["cost_entry_service"]
    entry = entries.create_manual_entry(
        project_id=project.id,
        command_id="d4-actual-1",
        description="Posted actual",
        amount=Decimal(25),
        currency_code=organization.base_currency,
        transaction_date=date(2026, 8, 5),
        cost_code_id=code.id,
    )
    entry = entries.submit(entry.id, expected_version=entry.row_version)
    _become_independent_decider(services)
    entries.approve(entry.id, expected_version=entry.row_version)
    _login(services, "admin", "ChangeMe123!")
    entry = entries.get_entry(entry.id)
    entry = entries.post(
        entry.id,
        expected_version=entry.row_version,
        posting_date=date(2026, 8, 5),
    )
    return project, budget, forecast, entry, code


def test_snapshot_reconciles_approved_budget_forecast_and_posted_actual(services) -> None:
    project, budget, forecast, _entry, _code = _approved_controls(services)

    snapshot = services["finance_service"].get_finance_snapshot(
        project.id, as_of=date(2026, 8, 31)
    )

    assert snapshot.budget == Decimal(100)
    assert snapshot.actual == Decimal(25)
    assert snapshot.forecast_etc == Decimal(80)
    assert snapshot.estimate_at_completion == Decimal(105)
    assert snapshot.budget_headroom == Decimal(-5)
    assert snapshot.approved_budget_id == budget.id
    assert snapshot.approved_forecast_id == forecast.id
    assert snapshot.approved_forecast_revision == forecast.revision
    assert snapshot.currency_basis == "PROJECT_CURRENCY"
    assert snapshot.period_granularity == "month"
    assert snapshot.reconciliation.is_reconciled is True
    assert snapshot.reconciliation.posted_actual_delta == Decimal(0)
    assert snapshot.reconciliation.open_commitment_delta == Decimal(0)
    assert snapshot.reconciliation.forecast_etc_delta == Decimal(0)
    assert sum(
        (row.amount for row in snapshot.ledger if row.stage == "forecast"),
        start=Decimal(0),
    ) == snapshot.forecast_etc
    assert [row.period_key for row in snapshot.cost_phasing] == ["2026-08"]
    assert snapshot.cost_phasing[0].actual == Decimal(25)
    assert snapshot.cost_phasing[0].forecast == Decimal(80)


def test_finance_excel_export_has_bounded_lineage_and_control_parity(
    services,
    tmp_path,
) -> None:
    project, budget, forecast, _entry, code = _approved_controls(services)
    output = tmp_path / "finance-d5.xlsx"

    reporting_api.generate_excel_report(
        services["reporting_service"],
        project.id,
        output,
        finance_service=services["finance_service"],
        as_of=date(2026, 8, 31),
        finance_ledger_offset=1,
        finance_ledger_limit=1,
    )

    workbook = load_workbook(output, data_only=True)
    finance = workbook["Finance"]
    metadata = {
        finance.cell(row=row, column=1).value: finance.cell(row=row, column=2).value
        for row in range(1, finance.max_row + 1)
    }
    assert metadata["Snapshot as of"] == "2026-08-31"
    assert metadata["Currency basis"].startswith("PROJECT_CURRENCY:")
    assert metadata["Approved budget version"] == f"{budget.id} / revision {budget.revision}"
    assert metadata["Approved forecast version"] == f"{forecast.id} / revision {forecast.revision}"
    assert metadata["Reconciliation status"] == "Reconciled"
    assert metadata["Ledger page limit"] == 1
    canonical = services["finance_performance_query"].get_cost_phasing(
        project.id,
        date_from=date(2026, 1, 1),
        date_to=date(2026, 8, 31),
        as_of_date=date(2026, 8, 31),
    )
    exported_periods = {
        finance.cell(row=row, column=1).value: row
        for row in range(1, finance.max_row + 1)
        if isinstance(finance.cell(row=row, column=1).value, str)
        and finance.cell(row=row, column=1).value.startswith("2026-")
    }
    assert set(exported_periods) == {item.period_key for item in canonical.periods}
    for period in canonical.periods:
        row = exported_periods[period.period_key]
        assert Decimal(str(finance.cell(row=row, column=2).value)) == period.planned_cost
        assert Decimal(str(finance.cell(row=row, column=3).value)) == period.open_commitment
        assert Decimal(str(finance.cell(row=row, column=4).value)) == period.posted_actual
        assert Decimal(str(finance.cell(row=row, column=5).value)) == period.forecast_cost
    for availability in canonical.series_availability:
        assert availability.availability in metadata[f"Cost phasing {availability.series_code}"]

    ledger = workbook["Finance Ledger"]
    headers = [cell.value for cell in ledger[1]]
    assert headers == [
        "Date",
        "Period Start",
        "Period End",
        "Source",
        "Source Type",
        "Stage",
        "Cost Type",
        "Cost Code ID",
        "Financial Period ID",
        "Reference Type",
        "Reference ID",
        "Reference",
        "Task ID",
        "Task",
        "Resource ID",
        "Resource",
        "Amount",
        "Currency",
    ]
    assert ledger.max_row == 2
    assert ledger.cell(row=2, column=8).value == code.id
    assert ledger.cell(row=2, column=9).value


def test_finance_pdf_export_uses_the_same_canonical_read_basis(
    services,
    tmp_path,
) -> None:
    project, _budget, _forecast, _entry, _code = _approved_controls(services)
    output = tmp_path / "finance-d5.pdf"

    reporting_api.generate_pdf_report(
        services["reporting_service"],
        project.id,
        output,
        temp_dir=tmp_path / "report-temp",
        finance_service=services["finance_service"],
        as_of=date(2026, 8, 31),
        finance_ledger_limit=1,
    )

    assert output.read_bytes().startswith(b"%PDF")


def test_snapshot_has_no_eac_or_vac_before_the_approved_forecast_basis(services) -> None:
    project, _budget, _forecast, _entry, _code = _approved_controls(services)

    snapshot = services["finance_service"].get_finance_snapshot(
        project.id, as_of=date(2026, 7, 31)
    )

    assert snapshot.actual == Decimal(0)
    assert snapshot.forecast_etc is None
    assert snapshot.estimate_at_completion is None
    assert snapshot.budget_headroom is None
    assert snapshot.approved_forecast_id is None


def test_posted_reversal_nets_actual_without_rewriting_forecast(services) -> None:
    project, _budget, forecast, entry, _code = _approved_controls(services)
    services["financial_period_service"].create_period(
        code="D4-2026-09",
        name="September 2026",
        fiscal_year=2026,
        period_number=9,
        start_date=date(2026, 9, 1),
        end_date=date(2026, 9, 30),
    )
    entries = services["cost_entry_service"]
    entries.reverse(
        entry.id,
        expected_version=entry.row_version,
        command_id="d4-reversal-1",
        posting_date=date(2026, 9, 20),
        reason="Correct the test posting",
    )

    snapshot = services["finance_service"].get_finance_snapshot(
        project.id, as_of=date(2026, 9, 30)
    )

    assert snapshot.actual == Decimal(0)
    assert snapshot.forecast_etc == Decimal(80)
    assert snapshot.estimate_at_completion == Decimal(80)
    assert snapshot.approved_forecast_id == forecast.id
    periods = {row.period_key: row.actual for row in snapshot.cost_phasing}
    assert periods["2026-08"] == Decimal(25)
    assert periods["2026-09"] == Decimal(-25)
    assert sum(periods.values(), Decimal(0)) == Decimal(0)


def test_as_of_selects_superseded_approved_forecast_version(services) -> None:
    project, _budget, first, _entry, code = _approved_controls(services)
    forecasts = services["forecast_version_service"]
    second = forecasts.create_forecast(
        project.id,
        name="September ETC",
        as_of_date=date(2026, 9, 1),
        generation_mode=ForecastGenerationMode.MANUAL,
        created_by="admin",
    )
    forecasts.add_line(
        second.id,
        cost_code_id=code.id,
        description="Revised remaining delivery",
        amount=Decimal(60),
        source_kind=ForecastLineSourceKind.MANUAL,
        source_type=ForecastLineSourceType.MANUAL_ESTIMATE,
        created_by="admin",
        expected_forecast_version=second.row_version,
    )
    second = forecasts.get_forecast(second.id)
    second = forecasts.submit_forecast(
        second.id, submitted_by="admin", expected_version=second.row_version
    )
    second = forecasts.approve_forecast(
        second.id, approved_by="admin", expected_version=second.row_version
    )

    august = services["finance_service"].get_finance_snapshot(
        project.id, as_of=date(2026, 8, 31)
    )
    september = services["finance_service"].get_finance_snapshot(
        project.id, as_of=date(2026, 9, 30)
    )

    assert august.approved_forecast_id == first.id
    assert august.forecast_etc == Decimal(80)
    assert sum((row.forecast for row in august.cost_phasing), Decimal(0)) == Decimal(80)
    assert september.approved_forecast_id == second.id
    assert september.forecast_etc == Decimal(60)
    assert next(item for item in september.cost_phasing_availability if item.series_code == "forecast").unphased_amount == Decimal(60)


def _clone_posted_entry(session, original_id: str, *, amount: str, posted_on: date, status: str = "posted") -> None:
    original = session.get(ProjectCostEntryORM, original_id)
    values = {column.key: getattr(original, column.key) for column in ProjectCostEntryORM.__table__.columns}
    identifier = str(uuid4())
    values.update(
        id=identifier,
        source_id=identifier,
        idempotency_key=identifier,
        amount=Decimal(amount),
        base_amount=Decimal(amount),
        transaction_date=posted_on,
        posting_date=posted_on,
        status=status,
        reverses_entry_id=None,
        reversed_by_entry_id=None,
    )
    if status not in {"posted", "reversed"}:
        for key in (
            "base_amount", "base_currency_code", "exchange_rate", "exchange_rate_date",
            "exchange_rate_source", "exchange_rate_captured_at", "posting_date",
            "financial_period_id", "posted_by", "posted_at",
        ):
            values[key] = None
    session.add(ProjectCostEntryORM(**values))


def test_r6e_cost_phasing_boundaries_reconcile_exactly_and_preserve_as_of(services) -> None:
    project, _budget, forecast, entry, _code = _approved_controls(services)
    session = services["session"]
    for amount, posted_on in (
        ("0.10", date(2026, 8, 1)),
        ("0.20", date(2026, 8, 31)),
        ("0.30", date(2026, 12, 31)),
        ("0.40", date(2027, 1, 1)),
        ("0.50", date(2026, 9, 1)),
        ("0.60", date(2026, 9, 15)),
    ):
        _clone_posted_entry(session, entry.id, amount=amount, posted_on=posted_on)
    for status in ("draft", "submitted", "approved"):
        _clone_posted_entry(session, entry.id, amount="99", posted_on=date(2026, 8, 31), status=status)

    original_line = session.scalars(
        select(ForecastLineORM).where(ForecastLineORM.forecast_id == forecast.id)
    ).one()
    line_values = {column.key: getattr(original_line, column.key) for column in ForecastLineORM.__table__.columns}
    for amount, starts, ends in (
        ("0.10", date(2026, 9, 1), date(2026, 9, 30)),
        ("0.20", None, None),
        ("0.30", date(2026, 12, 31), date(2027, 1, 1)),
    ):
        session.add(ForecastLineORM(**{
            **line_values,
            "id": str(uuid4()),
            "amount": Decimal(amount),
            "period_start": starts,
            "period_end": ends,
        }))
    session.commit()

    query = services["finance_performance_query"]
    facts = query.get_cost_phasing(
        project.id,
        date_from=date(2026, 8, 1),
        date_to=date(2027, 1, 31),
        as_of_date=date(2026, 12, 31),
    )
    periods = {row.period_key: row for row in facts.periods}
    assert periods["2026-08"].posted_actual == Decimal("25.30")
    assert periods["2026-09"].posted_actual == Decimal("1.10")
    assert periods["2026-12"].posted_actual == Decimal("0.30")
    assert "2027-01" not in periods
    assert sum((row.posted_actual for row in facts.periods), Decimal(0)) == Decimal("26.70")
    assert sum((row.forecast_cost for row in facts.periods), Decimal(0)) == Decimal("80.10")
    forecast_availability = next(item for item in facts.series_availability if item.series_code == "forecast")
    assert forecast_availability.phased_amount == Decimal("80.10")
    assert forecast_availability.unphased_amount == Decimal("0.50")
    assert forecast_availability.phased_amount + forecast_availability.unphased_amount == Decimal("80.60")

    snapshot = services["finance_service"].get_finance_snapshot(project.id, as_of=date(2026, 12, 31))
    snapshot_periods = {row.period_key: row for row in snapshot.cost_phasing}
    for key in ("2026-08", "2026-09", "2026-12"):
        assert snapshot_periods[key].actual == periods[key].posted_actual
        assert snapshot_periods[key].forecast == periods[key].forecast_cost
    assert snapshot.cost_phasing_availability == facts.series_availability

    mid_month = query.get_cost_phasing(
        project.id,
        date_from=date(2026, 9, 13),
        date_to=date(2026, 9, 30),
        as_of_date=date(2026, 9, 30),
    )
    assert [(row.period_key, row.posted_actual) for row in mid_month.periods] == [
        ("2026-09", Decimal("0.60"))
    ]


def test_cost_phasing_historical_actual_does_not_query_current_rates(services, monkeypatch) -> None:
    project, _budget, _forecast, entry, _code = _approved_controls(services)
    resource = services["resource_service"].create_resource("Historical cost resource")
    rate_service = services["rate_card_service"]
    card = rate_service.create_rate_card(name="Historical cost rates", project_id=project.id)
    rate_service.create_line(
        card.id,
        rate_type=RateType.COST,
        unit="HOUR",
        rate_amount=Decimal("123.45"),
        rate_currency=services["tenant_context_service"].get_active_organization().base_currency,
        resource_id=resource.id,
    )
    services["session"].get(ProjectCostEntryORM, entry.id).resource_id = resource.id
    services["session"].commit()
    query = services["finance_performance_query"]
    before = query.get_cost_phasing(
        project.id,
        date_from=date(2026, 8, 1),
        date_to=date(2026, 8, 31),
    )
    rate_service.deactivate_rate_card(card.id, expected_version=card.version)
    resolver = services["finance_service"]._labor._rate_resolver
    monkeypatch.setattr(
        resolver,
        "resolve_many",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("Current rate lookup")),
    )
    after = query.get_cost_phasing(
        project.id,
        date_from=date(2026, 8, 1),
        date_to=date(2026, 8, 31),
    )
    assert after.periods == before.periods


def test_r6e_integrated_performance_uses_one_approved_financial_basis(services) -> None:
    project, _budget, _forecast, _entry, _code = _approved_controls(services)
    tasks = services["task_service"]
    first = tasks.create_task(project.id, "Baseline task one", start_date=date(2026, 8, 3), duration_days=3)
    second = tasks.create_task(project.id, "Baseline task two", start_date=date(2026, 8, 6), duration_days=2)
    baselines = services["baseline_service"]
    baseline = baselines.create_baseline(project.id, "Cost-loaded control", rate_as_of=date(2026, 8, 1))
    session = services["session"]
    rows = session.execute(
        select(BaselineTaskORM).where(BaselineTaskORM.baseline_id == baseline.id)
    ).scalars().all()
    assert {row.task_id for row in rows} == {first.id, second.id}
    for row in rows:
        row.baseline_planned_cost = Decimal(60 if row.task_id == first.id else 40)
    session.commit()
    baselines.submit_baseline(baseline.id, submitted_by="admin")
    baselines.approve_baseline(baseline.id, approved_by="admin")
    with pytest.raises(NotFoundError):
        baselines.list_variance_records(baseline.id, expected_project_id="other-project")
    tasks.update_progress(first.id, percent_complete=50)
    tasks.update_progress(second.id, percent_complete=25)

    as_of = date(2026, 8, 31)
    performance = services["finance_performance_query"]
    evm = performance.get_evm(project.id, as_of_date=as_of)
    variance = performance.get_variance(project.id, as_of_date=as_of)
    phasing = performance.get_cost_phasing(
        project.id, date_from=date(2026, 8, 1), date_to=as_of, as_of_date=as_of
    )
    snapshot = services["finance_service"].get_finance_snapshot(project.id, as_of=as_of)
    report_evm = services["reporting_service"].get_earned_value(project.id, as_of=as_of)

    assert evm.availability == "available"
    assert (evm.bac, evm.pv, evm.ev, evm.ac, evm.etc, evm.eac, evm.vac) == (
        Decimal(100), Decimal(100), Decimal(40), Decimal(25), Decimal(80),
        Decimal(105), Decimal(-5),
    )
    assert evm.baseline_id == baseline.id
    assert (evm.cpi, evm.spi) == (Decimal("1.6"), Decimal("0.4"))
    assert report_evm.BAC == evm.bac
    assert report_evm.AC == evm.ac
    assert report_evm.EAC == evm.eac
    metrics = {item.metric_code: item for item in variance.metrics}
    assert metrics["cost_variance"].value == evm.ev - evm.ac == Decimal(15)
    assert metrics["schedule_variance"].value == evm.ev - evm.pv == Decimal(-60)
    assert metrics["vac"].value == evm.bac - evm.eac == Decimal(-5)
    assert metrics["budget_pressure"].value == evm.eac - snapshot.budget == Decimal(5)
    assert snapshot.budget_headroom == -metrics["budget_pressure"].value == Decimal(-5)
    assert sum((row.planned_cost for row in phasing.periods), Decimal(0)) == evm.bac
    assert sum((row.posted_actual for row in phasing.periods), Decimal(0)) == evm.ac
    assert sum((row.forecast_cost for row in phasing.periods), Decimal(0)) == evm.etc
    assert snapshot.cost_phasing[0].actual == phasing.periods[0].posted_actual
    assert snapshot.cost_phasing[0].planned == phasing.periods[0].planned_cost
